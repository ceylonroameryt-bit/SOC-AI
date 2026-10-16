import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_simple_box(doc, text, title=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F9FF")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="0284C7"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    if title:
        r_title = p.add_run(f"{title}\n")
        r_title.bold = True
        r_title.font.name = "Calibri"
        r_title.font.size = Pt(11)
        r_title.font.color.rgb = RGBColor(2, 132, 199)
        
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def generate_simple_proposal():
    doc = Document()

    # Standard clean margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Title & Subtitle
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("SOC-AI Platform")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(26)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(14, 116, 144)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("A Simple AI-Powered Cyber Threat Intelligence Dashboard | Project Proposal")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(13)
    r_sub.bold = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # 1. Project Overview
    h1 = doc.add_paragraph("1. Project Overview")
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    h1.runs[0].font.name = "Calibri"
    h1.runs[0].font.size = Pt(14)
    h1.runs[0].bold = True
    h1.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    p_desc = doc.add_paragraph(
        "SOC-AI is a clean, practical web dashboard built to help individuals and small security teams stay on top of daily cyber security news. "
        "Instead of checking dozens of blogs and websites manually, the app automatically collects security alerts in one place, "
        "removes duplicate stories, and uses AI to summarize long technical articles into quick, easy-to-read takeaways."
    )
    p_desc.paragraph_format.line_spacing = 1.15
    p_desc.paragraph_format.space_after = Pt(6)

    # 2. Why We Built It
    h2 = doc.add_paragraph("2. The Problem: Keeping Up with Cyber News is Hard")
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)
    h2.runs[0].font.name = "Calibri"
    h2.runs[0].font.size = Pt(14)
    h2.runs[0].bold = True
    h2.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    problems = [
        ("Too Many Websites to Check", "Cyber threats are reported across dozens of different websites, vendor blogs, and government portals. Manually checking them all every day is exhausting and takes too much time."),
        ("Lots of Duplicate Articles", "When a major security bug is discovered, 10 to 20 different websites write about the exact same thing. Readers waste time reading identical stories."),
        ("Too Technical and Long", "Most security advisories are filled with dense engineering jargon. Non-technical users and busy managers just need to know: 'What happened, and should we be worried?'")
    ]

    for title, desc in problems:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        rt = p.add_run(f"• {title}: ")
        rt.bold = True
        rt.font.name = "Calibri"
        rt.font.size = Pt(10)
        rt.font.color.rgb = RGBColor(185, 28, 28)
        rd = p.add_run(desc)
        rd.font.name = "Calibri"
        rd.font.size = Pt(10)
        rd.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 3. How the App Works
    h3 = doc.add_paragraph("3. How It Works (In 3 Simple Steps)")
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)
    h3.runs[0].font.name = "Calibri"
    h3.runs[0].font.size = Pt(14)
    h3.runs[0].bold = True
    h3.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    tbl = doc.add_table(rows=4, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Step", "What the System Does", "Why It Helps"]
    for j, val in enumerate(headers):
        cell = tbl.cell(0, j)
        set_cell_background(cell, "0E7490")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(val)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    steps_data = [
        ("Step 1: Collect", "Automatically reads RSS and news feeds from verified cyber security sites.", "All news in one place—no manual web browsing required."),
        ("Step 2: Clean", "Smart filters spot copycat stories and keep only unique updates.", "Cuts out repetitive clutter so you only read each story once."),
        ("Step 3: Summarize", "AI creates a 2-minute daily brief and highlights severity (Low/High).", "Fast, plain-English understanding without technical jargon.")
    ]

    for i, row_data in enumerate(steps_data):
        row = tbl.rows[i + 1]
        bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 70, 70, 90, 90)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            if j == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(14, 116, 144)
            else:
                r.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 4. Main Features
    h4 = doc.add_paragraph("4. Main Features")
    h4.paragraph_format.space_before = Pt(10)
    h4.paragraph_format.space_after = Pt(4)
    h4.runs[0].font.name = "Calibri"
    h4.runs[0].font.size = Pt(14)
    h4.runs[0].bold = True
    h4.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    features = [
        ("Live Threat Feed", "A clear scrolling list of the latest articles, labeled by severity (Low, Medium, High, Critical)."),
        ("AI Daily Briefing", "A 2-minute morning summary highlighting the most important cyber news of the day."),
        ("Extracted Indicators", "Pulls out useful clues like suspicious IP addresses, web domains, and vulnerability IDs (CVEs)."),
        ("Easy Categories & Search", "Quickly filter news by topic: Ransomware, Malware, Data Breaches, or Vulnerabilities.")
    ]

    for title, desc in features:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        rt = p.add_run(f"✔  {title}: ")
        rt.bold = True
        rt.font.name = "Calibri"
        rt.font.size = Pt(10)
        rt.font.color.rgb = RGBColor(16, 185, 129)
        rd = p.add_run(desc)
        rd.font.name = "Calibri"
        rd.font.size = Pt(10)
        rd.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 5. Technology Used
    h5 = doc.add_paragraph("5. Simple Technology Overview")
    h5.paragraph_format.space_before = Pt(10)
    h5.paragraph_format.space_after = Pt(4)
    h5.runs[0].font.name = "Calibri"
    h5.runs[0].font.size = Pt(14)
    h5.runs[0].bold = True
    h5.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    p_tech = doc.add_paragraph(
        "• Frontend: React & Tailwind CSS — Provides a clean, modern, dark-themed dashboard that works smoothly on any browser.\n"
        "• Backend Server: Node.js & Express — Fetches the news feeds, cleans the data, and runs the application logic.\n"
        "• AI Helper: Python (FastAPI) — Handles grouping similar news together and generating quick summaries.\n"
        "• Database: PostgreSQL — Stores the articles and ensures no duplicate stories clutter the dashboard."
    )
    p_tech.paragraph_format.line_spacing = 1.15
    p_tech.paragraph_format.space_after = Pt(8)

    # 6. Conclusion
    h6 = doc.add_paragraph("6. Conclusion")
    h6.paragraph_format.space_before = Pt(10)
    h6.paragraph_format.space_after = Pt(4)
    h6.runs[0].font.name = "Calibri"
    h6.runs[0].font.size = Pt(14)
    h6.runs[0].bold = True
    h6.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    add_simple_box(
        doc,
        "SOC-AI is a straightforward, practical tool that turns overwhelming daily cyber news into clean, "
        "easy-to-read alerts. It saves valuable time every morning, eliminates repetitive clutter, and gives everyone "
        "a clear picture of today's cybersecurity landscape.",
        title="IN SUMMARY"
    )

    out_file = "Simple_Project_Proposal.docx"
    doc.save(out_file)
    print(f"Simple proposal saved to: {out_file}")

if __name__ == "__main__":
    generate_simple_proposal()
