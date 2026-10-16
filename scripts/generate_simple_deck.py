from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_simple_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Clean, simple theme
    BG_COLOR = RGBColor(248, 250, 252)    # Light gray/white #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)     # Pure White
    CARD_BORDER = RGBColor(226, 232, 240) # Subtle gray #E2E8F0
    PRIMARY = RGBColor(14, 116, 144)      # Teal/Cyan #0E7490
    DARK_TEXT = RGBColor(15, 23, 42)      # Slate Dark #0F172A
    MUTED_TEXT = RGBColor(71, 85, 105)    # Slate Muted #475569
    ACCENT_GREEN = RGBColor(16, 185, 129) # Emerald #10B981
    ACCENT_RED = RGBColor(239, 68, 68)    # Coral Red #EF4444

    blank_layout = prs.slide_layouts[6]

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="PROJECT OVERVIEW"):
        cat_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(0.35))
        tf_cat = cat_box.text_frame
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = PRIMARY

        title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.95), Inches(11.3), Inches(0.7))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = DARK_TEXT

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Simple & Clean)
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(3.0))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = "SOC-AI Platform"
    p1.font.size = Pt(42)
    p1.font.bold = True
    p1.font.color.rgb = PRIMARY

    p2 = tf1.add_paragraph()
    p2.text = "A Simple AI-Powered Cyber Threat Dashboard"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = DARK_TEXT
    p2.space_before = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "Collecting, filtering, and summarizing daily cybersecurity news and alerts in one easy place."
    p3.font.size = Pt(15)
    p3.font.color.rgb = MUTED_TEXT
    p3.space_before = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 2: The Problem (Why we made this)
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "The Problem: Keeping Up with Cyber News is Hard", "WHY WE BUILT THIS")

    cards_s2 = [
        ("Too Many Websites", "Security news is spread across dozens of blogs and websites. Checking them manually every day takes too much time."),
        ("Duplicate Articles", "Multiple news sites report the exact same story. You end up reading the same thing over and over."),
        ("Too Technical to Skim", "Long articles are filled with technical details. It's hard to quickly know if an alert actually matters to you.")
    ]

    for i, (title, desc) in enumerate(cards_s2):
        x = Inches(1.0 + i * 3.85)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(3.6), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tb = s2.shapes.add_textbox(x + Inches(0.3), Inches(2.3), Inches(3.0), Inches(3.8))
        tf = tb.text_frame
        tf.word_wrap = True

        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(18)
        pt.font.bold = True
        pt.font.color.rgb = ACCENT_RED

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(13)
        pd.font.color.rgb = MUTED_TEXT
        pd.space_before = Pt(12)

    # -------------------------------------------------------------
    # SLIDE 3: The Solution (What this app does)
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "The Solution: A Single Smart Dashboard", "WHAT IT DOES")

    steps = [
        ("1. Collects Feeds", "Automatically grabs the latest security news and warnings from trusted cyber sources continuously."),
        ("2. Removes Duplicates", "Smart filters detect duplicate stories and show only unique, fresh updates."),
        ("3. AI Quick Summaries", "AI reads the long articles and creates a short, easy-to-read summary in seconds.")
    ]

    for i, (title, desc) in enumerate(steps):
        x = Inches(1.0 + i * 3.85)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(3.6), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tb = s3.shapes.add_textbox(x + Inches(0.3), Inches(2.3), Inches(3.0), Inches(3.8))
        tf = tb.text_frame
        tf.word_wrap = True

        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(18)
        pt.font.bold = True
        pt.font.color.rgb = PRIMARY

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(13)
        pd.font.color.rgb = MUTED_TEXT
        pd.space_before = Pt(12)

    # -------------------------------------------------------------
    # SLIDE 4: Main Features
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "Main Features of the App", "FEATURES")

    features = [
        ("Clean Live Feed", "See all fresh security news in one place, labeled by severity (Low, Medium, High, Critical)."),
        ("AI Daily Briefing", "Get a 2-minute daily morning recap of the top cyber events without reading long blogs."),
        ("Important Clues (IOCs)", "Automatically pulls out suspicious IP addresses, web domains, and vulnerability IDs (CVEs)."),
        ("Easy Search & Filter", "Filter by topic: Ransomware, Malware, Data Breaches, or Vulnerabilities.")
    ]

    for i, (title, desc) in enumerate(features):
        col = i % 2
        row = i // 2
        x = Inches(1.0 + col * 5.75)
        y = Inches(2.0 + row * 2.4)
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.5), Inches(2.1))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tb = s4.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), Inches(4.9), Inches(1.7))
        tf = tb.text_frame
        tf.word_wrap = True

        pt = tf.paragraphs[0]
        pt.text = f"✔  {title}"
        pt.font.size = Pt(16)
        pt.font.bold = True
        pt.font.color.rgb = PRIMARY

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(12.5)
        pd.font.color.rgb = MUTED_TEXT
        pd.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 5: Simple Technology Stack
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "How It Was Built (Simple Overview)", "TECHNOLOGY")

    tech_items = [
        ("Web Frontend", "Built with React and Tailwind CSS for a modern, responsive, and dark-themed user interface."),
        ("Backend Server", "Powered by Node.js and Express to fetch feeds, clean the data, and handle requests."),
        ("AI Microservice", "Lightweight Python service to group similar stories and generate quick summaries."),
        ("Database", "PostgreSQL database to store news, prevent duplicates, and save threat indicators.")
    ]

    for i, (title, desc) in enumerate(tech_items):
        col = i % 2
        row = i // 2
        x = Inches(1.0 + col * 5.75)
        y = Inches(2.0 + row * 2.4)
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.5), Inches(2.1))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tb = s5.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), Inches(4.9), Inches(1.7))
        tf = tb.text_frame
        tf.word_wrap = True

        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(16)
        pt.font.bold = True
        pt.font.color.rgb = DARK_TEXT

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(12.5)
        pd.font.color.rgb = MUTED_TEXT
        pd.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 6: Summary & Live Demo
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "Summary & Live Demo", "WRAP UP")

    c_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.0), Inches(10.3), Inches(4.5))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = CARD_BG
    c_box.line.color.rgb = PRIMARY

    tb_c = s6.shapes.add_textbox(Inches(2.0), Inches(2.4), Inches(9.3), Inches(3.7))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True

    pc1 = tf_c.paragraphs[0]
    pc1.text = "In Short: A Useful Tool to Save Time and Stay Safe"
    pc1.font.size = Pt(22)
    pc1.font.bold = True
    pc1.font.color.rgb = DARK_TEXT
    pc1.alignment = PP_ALIGN.CENTER

    points = [
        "Saves time: No need to browse 20+ websites every morning.",
        "Cuts the noise: Merges duplicate stories automatically.",
        "Clear answers: Quick AI summaries tell you what happened in plain English.",
        "Practical: Helps anyone stay aware of today's cyber threats."
    ]
    for p_text in points:
        p = tf_c.add_paragraph()
        p.text = f"✔  {p_text}"
        p.font.size = Pt(14)
        p.font.color.rgb = MUTED_TEXT
        p.space_before = Pt(10)
        p.alignment = PP_ALIGN.CENTER

    p_end = tf_c.add_paragraph()
    p_end.text = "Thank You! Let's take a look at the live app."
    p_end.font.size = Pt(17)
    p_end.font.bold = True
    p_end.font.color.rgb = PRIMARY
    p_end.space_before = Pt(18)
    p_end.alignment = PP_ALIGN.CENTER

    out_file = "Simple_Presentation.pptx"
    prs.save(out_file)
    print(f"Simple presentation saved to: {out_file}")

if __name__ == "__main__":
    create_simple_deck()
