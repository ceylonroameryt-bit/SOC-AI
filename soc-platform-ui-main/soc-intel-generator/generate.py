"""
SOC-INTEL Weekly Report Generator
==================================
Reads weekly_data.json → generates a professional Word (.docx) report.

Usage:
    python generate.py                   # reads weekly_data.json
    python generate.py --data my_data.json   # custom data file
    python generate.py --open            # open the file after generating

Author: Sujampathi Rathnayaka | NO ENTRY SOC Platform
"""

import json
import os
import sys
import argparse
import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ─── Colour constants ────────────────────────────────────────────────────────

C_INK       = "0F1117"
C_INK_LIGHT = "2D3142"
C_INK_MID   = "4A5068"
C_INK_SOFT  = "6B7280"
C_RULE_LIGHT = "F0F2F7"
C_RULE_MID  = "F8F9FC"
C_WHITE     = "FFFFFF"
C_RED       = "C0392B"
C_AMBER     = "B45309"
C_BLUE      = "1D4ED8"
C_GREY_BAR  = "D1D5DB"

SEVERITY_COLORS = {
    "critical":  C_RED,
    "high":      C_AMBER,
    "medium":    C_BLUE,
    "user_risk": C_AMBER,
    "low":       C_INK_SOFT,
}
SEVERITY_LABELS = {
    "critical":  "CRITICAL",
    "high":      "HIGH",
    "medium":    "MEDIUM",
    "user_risk": "USER RISK",
    "low":       "LOW",
}


# ─── XML / Style helpers ─────────────────────────────────────────────────────

def _set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_para_border(para, side="bottom", color=C_INK, size="4"):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    tag = OxmlElement(f"w:{side}")
    tag.set(qn("w:val"), "single")
    tag.set(qn("w:sz"), size)
    tag.set(qn("w:space"), "1")
    tag.set(qn("w:color"), color)
    pBdr.append(tag)
    pPr.append(pBdr)


def _run(para, text: str, bold=False, italic=False, size=10,
         color=C_INK, font="Inter"):
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = font
    run.font.color.rgb = RGBColor(*bytes.fromhex(color))
    return run


def _para(doc, text="", bold=False, italic=False, size=10,
          color=C_INK, space_before=0, space_after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    if text:
        _run(p, text, bold=bold, italic=italic, size=size, color=color)
    return p


def _section_heading(doc, num: str, title: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(8)
    _run(p, f"{num}  ·  {title}", bold=True, size=11, color=C_INK)
    _set_para_border(p, color=C_INK, size="4")
    return p


def _spacer(doc, pts=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(pts)
    return p


def _cell_para(cell, text="", bold=False, size=9, color=C_INK,
               space_before=4, space_after=4, align=None):
    p = cell.paragraphs[0] if not cell.paragraphs[0].text else cell.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    if text:
        _run(p, text, bold=bold, size=size, color=color)
    return p


def _add_cell_paragraph(cell, text="", bold=False, size=9, color=C_INK,
                         space_before=0, space_after=4):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        _run(p, text, bold=bold, size=size, color=color)
    return p


# ─── Document sections ───────────────────────────────────────────────────────

def build_header(doc, d: dict):
    meta = d["meta"]
    wk = meta["week_num"]
    dr = meta["date_range"]
    title = meta["title"]

    # Badge line
    badge = _para(doc, f"SOC-INTEL  /  WK{wk}  ·  {dr}",
                  size=8, color=C_INK_SOFT, bold=True,
                  space_before=0, space_after=4)

    # Main title
    h1 = _para(doc, title, bold=True, size=22, color=C_INK,
               space_before=0, space_after=6)

    # Subtitle
    st = d["stats"]
    subtitle_text = (
        f"Weekly cyber threat intelligence briefing — {dr}. "
        f"{st['articles_ingested']} ingested reports, {st['mitre_categorized']} "
        f"MITRE-categorized events, full ATT&CK tactic coverage."
    )
    _para(doc, subtitle_text, size=10, color=C_INK_SOFT,
          space_before=0, space_after=4)

    # Author line (with bottom border)
    author_text = (
        f"{meta['author']}  |  Security Analyst  |  "
        f"Published: {meta['published_date']}  |  NO ENTRY SOC Platform"
    )
    author_p = _para(doc, author_text, size=9, color=C_INK_MID,
                     space_before=0, space_after=8)
    _set_para_border(author_p, color=C_INK, size="12")

    _spacer(doc, 2)


def build_stat_strip(doc, d: dict):
    st = d["stats"]
    stats = [
        (str(st["articles_ingested"]),  "Threat articles ingested",   C_INK),
        (str(st["mitre_categorized"]),  "MITRE-categorized events",   C_BLUE),
        (str(st["tactic_coverage"]),    "ATT&CK tactic coverage",     C_INK),
        (str(st["active_zero_days"]),   "Active zero-days (KEV)",      C_RED),
    ]

    tbl = doc.add_table(rows=1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    row = tbl.rows[0]
    for i, (val, label, color) in enumerate(stats):
        cell = row.cells[i]
        _set_cell_bg(cell, C_RULE_LIGHT)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        p1 = cell.paragraphs[0]
        p1.paragraph_format.space_before = Pt(8)
        p1.paragraph_format.space_after = Pt(2)
        _run(p1, val, bold=True, size=20, color=color)

        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(8)
        _run(p2, label, size=8, color=C_INK_SOFT)

    _spacer(doc, 6)


def build_headlines(doc, d: dict):
    _section_heading(doc, "01", "Critical Headlines")

    for hl in d["headlines"]:
        severity = hl.get("severity", "medium")
        badge_label = SEVERITY_LABELS.get(severity, severity.upper())
        badge_color = SEVERITY_COLORS.get(severity, C_INK_SOFT)

        # Title + badge row (2-col table)
        t1 = doc.add_table(rows=1, cols=2)
        t1.alignment = WD_TABLE_ALIGNMENT.CENTER
        t1.style = "Table Grid"
        t1.columns[0].width = Inches(5.4)
        t1.columns[1].width = Inches(1.1)

        # Title cell
        c0 = t1.rows[0].cells[0]
        _set_cell_bg(c0, C_RULE_LIGHT)
        p = c0.paragraphs[0]
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(5)
        _run(p, hl["title"], bold=True, size=10, color=C_INK)

        # Badge cell
        c1 = t1.rows[0].cells[1]
        _set_cell_bg(c1, C_RULE_LIGHT)
        c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        pb = c1.paragraphs[0]
        pb.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pb.paragraph_format.space_before = Pt(5)
        pb.paragraph_format.space_after = Pt(5)
        _run(pb, badge_label, bold=True, size=8, color=badge_color)

        # Body table (1-col)
        t2 = doc.add_table(rows=1, cols=1)
        t2.alignment = WD_TABLE_ALIGNMENT.CENTER
        t2.style = "Table Grid"
        cb = t2.rows[0].cells[0]

        # Body text
        p2 = cb.paragraphs[0]
        p2.paragraph_format.space_before = Pt(6)
        p2.paragraph_format.space_after = Pt(4)
        _run(p2, hl.get("body", ""), size=9.5, color=C_INK_LIGHT)

        # CVEs
        for cve in hl.get("cves", []):
            pc = cb.add_paragraph()
            pc.paragraph_format.space_before = Pt(2)
            pc.paragraph_format.space_after = Pt(2)
            _run(pc, f"  {cve['id']}  ", bold=True, size=9, color=C_INK)
            _run(pc, f"  {cve['desc']}", size=9, color=C_INK_LIGHT)

        # Action
        if hl.get("action"):
            pa = cb.add_paragraph()
            pa.paragraph_format.space_before = Pt(6)
            pa.paragraph_format.space_after = Pt(6)
            _run(pa, "->  ", bold=True, size=9, color=C_INK)
            _run(pa, hl["action"], size=9, color=C_INK_MID)

        _spacer(doc, 4)


def build_tactic_volume(doc, d: dict):
    _section_heading(doc, "02", "Where the Volume Is")

    _para(doc,
          f"Across {d['stats']['mitre_categorized']} categorized events, three tactics account "
          f"for the overwhelming majority of observed activity.",
          size=10, color=C_INK_SOFT, space_after=8)

    vols = d["tactic_volume"]
    tbl = doc.add_table(rows=1, cols=len(vols))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    for i, v in enumerate(vols):
        cell = tbl.rows[0].cells[i]
        _set_cell_bg(cell, C_RULE_MID)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        p1 = cell.paragraphs[0]
        p1.paragraph_format.space_before = Pt(8)
        p1.paragraph_format.space_after = Pt(0)
        _run(p1, f"{v['rank']}  ", bold=True, size=14, color=C_GREY_BAR)
        _run(p1, str(v["count"]), bold=True, size=20, color=C_INK)

        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(8)
        _run(p2, v["tactic"], size=9, color=C_INK_SOFT)

    note_text = d.get("tactic_volume_note", "")
    if note_text:
        note = _para(doc, note_text, italic=True, size=9.5, color=C_INK_MID,
                     space_before=6, space_after=4)

    _spacer(doc, 2)


def build_techniques(doc, d: dict):
    _section_heading(doc, "03", "Hottest Technique per Tactic")

    tbl = doc.add_table(rows=1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    # Header
    headers = ["Tactic", "Technique", "ATT&CK ID", "Hits this week"]
    for i, h in enumerate(headers):
        c = tbl.rows[0].cells[i]
        _set_cell_bg(c, C_INK)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        _run(p, h, bold=True, size=8.5, color=C_WHITE)

    # Data rows
    for idx, t in enumerate(d["techniques"]):
        bg = C_WHITE if idx % 2 == 0 else C_RULE_LIGHT
        row = tbl.add_row()
        cells_data = [
            (t["tactic"],     True,  C_INK),
            (t["technique"],  False, C_INK_LIGHT),
            (t["attck_id"],   True,  C_BLUE),
            (str(t["hits"]),  True,  C_INK),
        ]
        for i, (text, bold, color) in enumerate(cells_data):
            c = row.cells[i]
            _set_cell_bg(c, bg)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            _run(p, text, bold=bold, size=9, color=color)

    _spacer(doc, 4)


def build_soc_pillars(doc, d: dict):
    _section_heading(doc, "04", "What High-Performing SOCs Are Doing Differently")

    intro = d.get("soc_pillars_intro", "")
    if intro:
        _para(doc, intro, size=10, color=C_INK_SOFT, space_after=8)

    pillars = d["soc_pillars"]
    rows_needed = (len(pillars) + 1) // 2
    tbl = doc.add_table(rows=rows_needed, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    for idx, p in enumerate(pillars):
        r, c = divmod(idx, 2)
        cell = tbl.rows[r].cells[c]
        _set_cell_bg(cell, C_RULE_MID)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

        p1 = cell.paragraphs[0]
        p1.paragraph_format.space_before = Pt(6)
        p1.paragraph_format.space_after = Pt(2)
        _run(p1, f"{p['num']}  ·  {p['label']}", bold=True, size=8, color=C_INK_SOFT)

        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(3)
        _run(p2, p["title"], bold=True, size=9.5, color=C_INK)

        p3 = cell.add_paragraph()
        p3.paragraph_format.space_before = Pt(0)
        p3.paragraph_format.space_after = Pt(6)
        _run(p3, p["body"], size=9, color=C_INK_SOFT)

    _spacer(doc, 4)


def build_breach_tracker(doc, d: dict):
    _section_heading(doc, "05", "Major Breach Tracker")

    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    for i, h in enumerate(["Organisation", "Impact", "Type"]):
        c = tbl.rows[0].cells[i]
        _set_cell_bg(c, C_INK)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        _run(p, h, bold=True, size=8.5, color=C_WHITE)

    for idx, b in enumerate(d["breaches"]):
        bg = C_WHITE if idx % 2 == 0 else C_RULE_LIGHT
        row = tbl.add_row()
        items = [(b["org"], True, C_INK), (b["impact"], False, C_INK_LIGHT), (b["type"], False, C_INK_LIGHT)]
        for i, (text, bold, color) in enumerate(items):
            c = row.cells[i]
            _set_cell_bg(c, bg)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            _run(p, text, bold=bold, size=9, color=color)

    note = d.get("breach_note", "")
    if note:
        _para(doc, note, italic=True, size=9.5, color=C_INK_MID,
              space_before=6, space_after=4)

    _spacer(doc, 4)


def build_actions(doc, d: dict):
    _section_heading(doc, "06", "This Weekend")

    for i, action in enumerate(d["actions"], 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        _run(p, f"{i}.  ", bold=True, size=10, color=C_INK)
        _run(p, action["title"] + "  ", bold=True, size=10, color=C_INK)
        _run(p, action["detail"], size=10, color=C_INK_MID)


def build_footer(doc, d: dict):
    meta = d["meta"]
    _spacer(doc, 8)
    sep = _para(doc, space_before=4, space_after=0)
    _set_para_border(sep, color=C_INK, size="12")

    footer_text = (
        f"Resources: CISA KEV Catalog  |  MITRE ATT&CK Framework  |  "
        f"Microsoft Security Response Center  |  NO ENTRY SOC Platform\n"
        f"Data: NO ENTRY SOC Platform  |  Week of {meta['date_range']}  |  "
        f"Mapped to MITRE ATT&CK Enterprise v15  |  "
        f"SOC-INTEL / WK{meta['week_num']}  |  {meta['author']}"
    )
    _para(doc, footer_text, size=8, color="9CA3AF", space_before=6, space_after=0)


# ─── Main builder ─────────────────────────────────────────────────────────────

def build_document(data: dict, output_path: Path):
    doc = Document()

    # Page layout
    sec = doc.sections[0]
    sec.page_width    = Cm(21)
    sec.page_height   = Cm(29.7)
    sec.left_margin   = Cm(2)
    sec.right_margin  = Cm(2)
    sec.top_margin    = Cm(1.8)
    sec.bottom_margin = Cm(1.8)

    # Default font
    doc.styles["Normal"].font.name = "Inter"
    doc.styles["Normal"].font.size = Pt(10)

    build_header(doc, data)
    build_stat_strip(doc, data)
    build_headlines(doc, data)
    build_tactic_volume(doc, data)
    build_techniques(doc, data)
    build_soc_pillars(doc, data)
    build_breach_tracker(doc, data)
    build_actions(doc, data)
    build_footer(doc, data)

    doc.save(str(output_path))
    return output_path


# ─── CLI ─────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="SOC-INTEL Weekly Report Generator — NO ENTRY SOC Platform"
    )
    parser.add_argument(
        "--data", default="weekly_data.json",
        help="Path to the weekly data JSON file (default: weekly_data.json)"
    )
    parser.add_argument(
        "--out", default=None,
        help="Output directory (default: ./reports/)"
    )
    parser.add_argument(
        "--open", action="store_true",
        help="Open the generated file after creation"
    )
    args = parser.parse_args()

    # Load data
    data_path = Path(args.data)
    if not data_path.exists():
        print(f"[ERROR] Data file not found: {data_path}")
        sys.exit(1)

    with open(data_path, encoding="utf-8") as f:
        data = json.load(f)

    # Output path
    out_dir = Path(args.out) if args.out else Path(__file__).parent / "reports"
    out_dir.mkdir(parents=True, exist_ok=True)

    wk  = data["meta"]["week_num"]
    year = datetime.date.today().year
    filename = f"SOC-INTEL-WK{wk:02d}-{year}.docx"
    output_path = out_dir / filename

    # Generate
    print(f"\n  SOC-INTEL Weekly Report Generator")
    print(f"  NO ENTRY SOC Platform | {data['meta']['author']}")
    print(f"  {'-'*44}")
    print(f"  Input  : {data_path.resolve()}")
    print(f"  Output : {output_path.resolve()}")
    print(f"  Week   : WK{wk:02d} | {data['meta']['date_range']}")
    print(f"\n  Generating...")

    result = build_document(data, output_path)

    size_kb = os.path.getsize(result) // 1024
    print(f"  Done!   {filename}  ({size_kb} KB)")
    print(f"\n  {'-'*44}")
    print(f"  To export PDF: Open in Word -> File -> Export -> Create PDF/XPS")
    print()

    if args.open:
        os.startfile(str(result))


if __name__ == "__main__":
    main()
